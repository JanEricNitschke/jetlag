"""Tests: per-season CSV/Excel export and the manual Google Sheet sync/pull."""

from __future__ import annotations

import csv
from pathlib import Path
from unittest import mock

import gspread
import pytest
from openpyxl import load_workbook

from jetlag_maps.cli import ExportArgs, PullArgs, SyncArgs
from jetlag_maps.config import GOOGLE_SHEET_ID
from jetlag_maps.format import SeasonFile, load_seasons
from jetlag_maps.sheets import (
    run_pull,
    run_sync,
    season_tab_name,
    season_tab_rows,
    season_tab_values,
)
from jetlag_maps.tabular import (
    LegRow,
    StopRow,
    csv_tables,
    export_tables,
    run_export,
    season_sheet_names,
)


def _read_rows[Row: StopRow | LegRow](path: Path, model: type[Row]) -> list[Row]:
    with path.open("r", newline="", encoding="utf-8") as stream:
        return [model.model_validate(record) for record in csv.DictReader(stream)]


def _sorted_seasons(committed_data: Path) -> list[SeasonFile]:
    return sorted(load_seasons(committed_data), key=lambda season: season.season)


def test_export_writes_only_per_season_csvs(
    committed_data: Path, tmp_path: Path
) -> None:
    """The CSV export writes exactly one stops/legs file pair per season."""
    run_export(ExportArgs(data=committed_data, out=tmp_path))
    for season in load_seasons(committed_data):
        stem = f"season_{season.season:02d}"
        stops = _read_rows(tmp_path / f"{stem}_stops.csv", StopRow)
        legs = _read_rows(tmp_path / f"{stem}_legs.csv", LegRow)
        assert (stops, legs) == export_tables([season])
    written = {path.name for path in tmp_path.iterdir()}
    expected = {
        f"season_{season.season:02d}_{table}.csv"
        for season in load_seasons(committed_data)
        for table in ("stops", "legs")
    }
    assert written == expected


def test_csv_tables_read_per_season_files(committed_data: Path, tmp_path: Path) -> None:
    """csv_tables re-reads the per-season files into the combined tables."""
    run_export(ExportArgs(data=committed_data, out=tmp_path))
    assert csv_tables(tmp_path) == export_tables(load_seasons(committed_data))


def test_xlsx_export_has_one_tab_per_season(
    committed_data: Path, tmp_path: Path
) -> None:
    """The Excel export holds one stops and one legs sheet per season."""
    out = tmp_path / "tables.xlsx"
    run_export(ExportArgs(data=committed_data, out=out))
    workbook = load_workbook(out)
    try:
        assert workbook.sheetnames == [
            name
            for season in _sorted_seasons(committed_data)
            for name in season_sheet_names(season)
        ]
    finally:
        workbook.close()


def test_season_tab_names_follow_plan_layout(committed_data: Path) -> None:
    """Tab titles match the "Season N - <Name>" layout of the plan."""
    names = {season_tab_name(season) for season in load_seasons(committed_data)}
    assert {
        "Season 0 - Crime Spree",
        "Season 1 - Connect 4",
        "Season 2 - Circumnavigation",
        "Season 3 - Tag Eur It",
        "Season 4 - Battle 4 America",
        "Season 5 - Race to the End",
        "Season 6 - Capture the Flag",
        "Season 7 - Tag Eur It 2",
        "Season 8 - Arctic Escape",
        "Season 9 - Hide + Seek",
    } <= names


def test_season_tab_values_contain_both_tables(
    committed_data: Path,
) -> None:
    """A tab holds the stops table, a blank row, then the legs table."""
    season = load_seasons(committed_data)[0]
    stops, legs = export_tables([season])
    values = season_tab_values(season)
    blank = [""] * len(StopRow.model_fields)
    assert values[0] == list(StopRow.model_fields)
    assert values[1] == list(stops[0].model_dump().values())
    assert values[1 + len(stops)] == blank
    assert values[1 + len(stops) + 1] == list(LegRow.model_fields)
    assert len(values) == len(stops) + len(legs) + 3


def test_season_tab_rows_parses_pushed_values(committed_data: Path) -> None:
    """The values pushed by sync parse back into the same table rows."""
    for season in load_seasons(committed_data):
        stops, legs = export_tables([season])
        values = season_tab_values(season)
        values.append([""] * len(StopRow.model_fields))  # trailing blank row
        assert season_tab_rows(values) == (stops, legs)


def test_season_tab_rows_skips_blank_and_leading_rows() -> None:
    """Blank rows and stray rows before the first header are skipped."""
    values = [
        ["a note"],
        list(StopRow.model_fields),
        [""] * len(StopRow.model_fields),
        list(LegRow.model_fields),
        [""] * len(LegRow.model_fields),
    ]
    assert season_tab_rows(values) == ([], [])


def test_run_sync_creates_missing_tabs(
    committed_data: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Sync adds one tab per season and pushes each season's values."""
    credentials = tmp_path / "service-account.json"
    credentials.write_text("{}", encoding="utf-8")
    service_account = mock.MagicMock()
    client = service_account.return_value
    spreadsheet = client.open_by_key.return_value
    spreadsheet.worksheet.side_effect = gspread.WorksheetNotFound
    monkeypatch.setattr(gspread, "service_account", service_account)

    run_sync(SyncArgs(data=committed_data, credentials=credentials))

    service_account.assert_called_once_with(filename=str(credentials))
    client.open_by_key.assert_called_once_with(GOOGLE_SHEET_ID)
    seasons = _sorted_seasons(committed_data)
    added = [call.kwargs["title"] for call in spreadsheet.add_worksheet.call_args_list]
    assert added == [season_tab_name(season) for season in seasons]
    worksheet = spreadsheet.add_worksheet.return_value
    pushed = [call.kwargs["values"] for call in worksheet.update.call_args_list]
    assert pushed == [season_tab_values(season) for season in seasons]
    assert worksheet.clear.call_count == len(seasons)


def test_run_sync_reuses_existing_tabs(
    committed_data: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Existing tabs are cleared and rewritten instead of re-created."""
    credentials = tmp_path / "service-account.json"
    credentials.write_text("{}", encoding="utf-8")
    service_account = mock.MagicMock()
    client = service_account.return_value
    spreadsheet = client.open_by_key.return_value
    monkeypatch.setattr(gspread, "service_account", service_account)

    run_sync(SyncArgs(data=committed_data, credentials=credentials))

    spreadsheet.add_worksheet.assert_not_called()
    worksheet = spreadsheet.worksheet.return_value
    titles = [call.args[0] for call in spreadsheet.worksheet.call_args_list]
    seasons = _sorted_seasons(committed_data)
    assert titles == [season_tab_name(season) for season in seasons]
    assert worksheet.clear.call_count == len(seasons)
    assert worksheet.update.call_count == len(seasons)


def test_run_sync_without_credentials_exits_cleanly(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A missing credentials file exits non-zero with an actionable hint."""
    service_account = mock.MagicMock()
    monkeypatch.setattr(gspread, "service_account", service_account)
    with pytest.raises(SystemExit) as excinfo:
        run_sync(SyncArgs(data=tmp_path, credentials=tmp_path / "missing.json"))
    message = excinfo.value.code
    assert isinstance(message, str)
    assert "Service-account credentials not found" in message
    service_account.assert_not_called()


def _oversized_leg_row(length: int) -> LegRow:
    """Build a legs row whose encoded geometry is ``length`` characters."""
    return LegRow(
        season=5,
        journey_team="Sam Denby",
        stop_index=0,
        leg_index=0,
        mode="car",
        note="",
        players_override="",
        geometry="0.1 " * (length // 4),
    )


def test_xlsx_export_exits_on_oversized_geometry(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """The Excel export refuses geometries over the cell-size limit."""
    oversized = [_oversized_leg_row(40000)]
    monkeypatch.setattr(
        "jetlag_maps.tabular.leg_rows",
        lambda _: oversized,  # pyrefly: ignore[implicit-any-lambda]
        raising=False,
    )
    with pytest.raises(SystemExit) as excinfo:
        run_export(ExportArgs(data=tmp_path, out=tmp_path / "tables.xlsx"))
    message = excinfo.value.code
    assert isinstance(message, str)
    assert "32767" in message
    assert "downsample" in message.lower()


def test_run_sync_exits_on_oversized_geometry(
    committed_data: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Sync refuses to push geometries over the cell-size limit."""
    oversized = [_oversized_leg_row(40000)]
    monkeypatch.setattr(
        "jetlag_maps.sheets.leg_rows",
        lambda _: oversized,  # pyrefly: ignore[implicit-any-lambda]
        raising=False,
    )
    credentials = tmp_path / "service-account.json"
    credentials.write_text("{}", encoding="utf-8")
    service_account = mock.MagicMock()
    monkeypatch.setattr(gspread, "service_account", service_account)
    with pytest.raises(SystemExit) as excinfo:
        run_sync(SyncArgs(data=committed_data, credentials=credentials))
    message = excinfo.value.code
    assert isinstance(message, str)
    assert "32767" in message
    assert "downsample" in message.lower()
    service_account.assert_not_called()


def _pull_spreadsheet(committed_data: Path) -> mock.MagicMock:
    """Build a mocked spreadsheet holding one tab per committed season."""
    spreadsheet = mock.MagicMock()
    tabs = []
    for season in _sorted_seasons(committed_data):
        tab = mock.MagicMock()
        tab.title = season_tab_name(season)
        tab.get_values.return_value = season_tab_values(season)
        tabs.append(tab)
    unrelated = mock.MagicMock()
    unrelated.title = "Notes"
    unrelated.get_values.return_value = [["not a season tab"]]
    tabs.append(unrelated)
    spreadsheet.worksheets.return_value = tabs
    return spreadsheet


def test_run_pull_writes_season_jsons(
    committed_data: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Pull reads every season tab and writes the season JSON files back."""
    credentials = tmp_path / "service-account.json"
    credentials.write_text("{}", encoding="utf-8")
    service_account = mock.MagicMock()
    service_account.return_value.open_by_key.return_value = _pull_spreadsheet(
        committed_data
    )
    monkeypatch.setattr(gspread, "service_account", service_account)

    out = tmp_path / "pulled"
    run_pull(PullArgs(out=out, credentials=credentials))

    service_account.assert_called_once_with(filename=str(credentials))
    service_account.return_value.open_by_key.assert_called_once_with(GOOGLE_SHEET_ID)
    assert load_seasons(out) == load_seasons(committed_data)


def test_run_pull_skips_non_season_tabs(
    committed_data: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Only tabs matching the "Season N - <Name>" layout are pulled."""
    credentials = tmp_path / "service-account.json"
    credentials.write_text("{}", encoding="utf-8")
    spreadsheet = _pull_spreadsheet(committed_data)
    service_account = mock.MagicMock()
    service_account.return_value.open_by_key.return_value = spreadsheet
    monkeypatch.setattr(gspread, "service_account", service_account)

    run_pull(PullArgs(out=tmp_path / "pulled", credentials=credentials))

    tabs = spreadsheet.worksheets.return_value
    tabs[-1].get_values.assert_not_called()  # the "Notes" tab
    for tab in tabs[:-1]:
        tab.get_values.assert_called_once()


def test_run_pull_without_credentials_exits_cleanly(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A missing credentials file exits non-zero with an actionable hint."""
    service_account = mock.MagicMock()
    monkeypatch.setattr(gspread, "service_account", service_account)
    with pytest.raises(SystemExit) as excinfo:
        run_pull(PullArgs(out=tmp_path, credentials=tmp_path / "missing.json"))
    message = excinfo.value.code
    assert isinstance(message, str)
    assert "Service-account credentials not found" in message
    service_account.assert_not_called()
